"""Audit-trail export helpers (PDF + Excel).

Used by GET /api/audit-logs/export in web_app.py.
Excel uses openpyxl (already in requirements.txt).
PDF uses reportlab (added to requirements.txt).
"""
from __future__ import annotations

from datetime import datetime, timezone
from io import BytesIO
from typing import List

EXPORT_HEADERS = [
    "Date/Time", "User", "Account Type", "Event",
    "Module", "Description", "IP Address", "Status",
]


def export_rows(records) -> List[list]:
    rows = []
    for r in records or []:
        rows.append([
            str(r.get("timestamp", "") or ""),
            str(r.get("email", "") or ""),
            str(r.get("accountType", "") or ""),
            str(r.get("event", "") or ""),
            str(r.get("module", "") or ""),
            str(r.get("description", "") or ""),
            str(r.get("ipAddress", "") or ""),
            str(r.get("status", "") or ""),
        ])
    return rows


def build_excel(rows: List[list], *, headers: List[str] | None = None, sheet_name: str = "Audit Trail") -> BytesIO:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = sheet_name
    ws.append(headers or EXPORT_HEADERS)
    header_fill = PatternFill(start_color="2A276E", end_color="2A276E", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True, size=10)
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    for row in rows:
        ws.append(row)
    width_headers = headers or EXPORT_HEADERS
    widths = [max(12, min(60, len(str(width_headers[index])) + 8)) for index in range(len(width_headers))]
    for idx, width in enumerate(widths, start=1):
        ws.column_dimensions[ws.cell(row=1, column=idx).column_letter].width = width
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def _pdf_text(value) -> str:
    """Escape audit values for ReportLab Paragraph markup."""
    from xml.sax.saxutils import escape
    text = escape(str(value or "-").strip())[:400]
    return text or "-"


def build_pdf(rows: List[list], *, headers: List[str] | None = None, title: str = "Audit Trail Report") -> BytesIO:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    generated = datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M UTC")
    buf = BytesIO()
    doc = SimpleDocTemplate(
        buf, pagesize=landscape(A4),
        leftMargin=10 * mm, rightMargin=10 * mm,
        topMargin=12 * mm, bottomMargin=12 * mm,
        title=title,
    )
    styles = getSampleStyleSheet()
    story = [
        Paragraph(title, styles["Title"]),
        Paragraph(
            f"Generated: {generated} &nbsp;&nbsp;|&nbsp;&nbsp; Records: {len(rows)}",
            styles["Normal"],
        ),
        Spacer(1, 6 * mm),
    ]
    cell_style = styles["BodyText"]
    cell_style.fontSize = 7
    cell_style.leading = 9
    header_style = styles["BodyText"].clone("ExportHeader")
    header_style.fontSize = 7
    header_style.leading = 9
    header_style.textColor = colors.white
    header_cells = [Paragraph(f"<b>{h}</b>", header_style) for h in (headers or EXPORT_HEADERS)]
    table_data = [header_cells]
    for row in rows:
        table_data.append([Paragraph(_pdf_text(c), cell_style) for c in row])
    available_width = landscape(A4)[0] - 20 * mm
    col_widths = [available_width / max(1, len(headers or EXPORT_HEADERS))] * len(headers or EXPORT_HEADERS)
    table = Table(table_data, colWidths=col_widths, repeatRows=1)
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#2A276E")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTSIZE", (0, 0), (-1, -1), 7),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#CCCCCC")),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F4F5FB")]),
    ]))
    story.append(table)
    doc.build(story)
    buf.seek(0)
    return buf
